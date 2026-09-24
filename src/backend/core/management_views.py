import json
import re
import uuid
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt

from .models import Employee, KnowledgeDocument


def employee_data(employee):
    return {'id': employee.pk, 'user_id': employee.user_id, 'full_name': employee.full_name,
            'email': employee.user.email, 'phone': employee.phone, 'active': employee.active,
            'role': employee.role, 'role_label': employee.get_role_display(),
            'created_at': employee.created_at.isoformat()}


@csrf_exempt
def employees(request, employee_id=None):
    if request.method == 'GET' and employee_id is None:
        return JsonResponse({'results': [employee_data(e) for e in Employee.objects.select_related('user').order_by('-created_at')]})
    if request.method not in ('POST', 'PATCH'):
        return JsonResponse({'message': 'Method not allowed.'}, status=405)
    try:
        data = json.loads(request.body)
        if not isinstance(data, dict):
            raise ValueError()
        with transaction.atomic():
            if request.method == 'PATCH':
                employee = Employee.objects.select_for_update().filter(pk=employee_id).first()
                if not employee:
                    return JsonResponse({'message': 'Không tìm thấy nhân viên.'}, status=404)
                if not isinstance(data.get('active'), bool):
                    raise ValueError()
                employee.active = data['active']
                employee.save(update_fields=['active'])
            else:
                name, email, phone = (str(data.get(key, '')).strip() for key in ('full_name', 'email', 'phone'))
                if not name or len(name) > 150 or not re.fullmatch(r'0\d{9}', phone):
                    raise ValueError()
                validate_email(email)
                role = data.get('role')
                if role not in Employee.Role.values:
                    return JsonResponse({'message': 'Chọn chức vụ Nhân viên kho hoặc Nhân viên xử lý.'}, status=400)
                user = User.objects.create_user(username=f'employee_{uuid.uuid4().hex}', email=email, first_name=name, password=None)
                employee = Employee.objects.create(user=user, full_name=name, phone=phone, role=role)
        return JsonResponse({'employee': employee_data(employee)}, status=201 if request.method == 'POST' else 200)
    except (ValueError, TypeError, ValidationError):
        return JsonResponse({'message': 'Nhập họ tên, email và số điện thoại hợp lệ.'}, status=400)
    except IntegrityError:
        return JsonResponse({'message': 'Số điện thoại nhân viên đã tồn tại.'}, status=400)


def document_data(document):
    return {'id': document.pk, 'name': document.name, 'size': document.size,
            'created_at': document.created_at.isoformat(), 'characters': len(document.text)}


@csrf_exempt
def documents(request, document_id=None):
    if request.method == 'GET':
        if document_id is None:
            return JsonResponse({'results': [document_data(d) for d in KnowledgeDocument.objects.defer('content').order_by('-created_at')]})
        document = KnowledgeDocument.objects.defer('content').filter(pk=document_id).first()
        if not document:
            return JsonResponse({'message': 'Không tìm thấy tài liệu.'}, status=404)
        return JsonResponse({'document': {**document_data(document), 'text': document.text}})
    if request.method != 'POST' or document_id is not None:
        return JsonResponse({'message': 'Method not allowed.'}, status=405)
    upload = request.FILES.get('file')
    if not upload or Path(upload.name).suffix.lower() not in ('.md', '.csv', '.pdf', '.docx'):
        return JsonResponse({'message': 'Chọn file MD, CSV, PDF hoặc Word (.docx).'}, status=400)
    if upload.size > 5 * 1024 * 1024:
        return JsonResponse({'message': 'Dung lượng tối đa 5 MB mỗi file.'}, status=400)
    raw = upload.read()
    try:
        extension = Path(upload.name).suffix.lower()
        if extension == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 300:
                raise ValueError()
            text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
        elif extension == '.docx':
            from docx import Document
            with ZipFile(BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError()
            document = Document(BytesIO(raw))
            text = '\n'.join([p.text for p in document.paragraphs] +
                             [' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
        else:
            text = raw.decode('utf-8-sig')
        if not text.strip() or '\x00' in text:
            raise ValueError()
        if len(text) > 2_000_000:
            raise ValueError()
    except Exception:
        return JsonResponse({'message': 'Không đọc được nội dung. MD/CSV cần UTF-8; PDF cần có lớp văn bản, không khóa mật khẩu (tối đa 300 trang); Word dùng .docx.'}, status=400)
    document = KnowledgeDocument.objects.create(name=upload.name[:255], content=raw, text=text, size=len(raw))
    return JsonResponse({'document': document_data(document)}, status=201)
